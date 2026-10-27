"""
cover_build.py - Dựng phần bìa và phần đầu báo cáo BTL (mẫu TNUT) ngay trong cùng
tài liệu Word, tránh mọi lỗi do ghép hai tệp riêng.

Gồm ba trang: bìa chính, bìa lót và nhiệm vụ bài tập lớn. Phần bìa nằm ở section
riêng KHÔNG đánh số trang; thân báo cáo ở section sau bắt đầu lại từ số 1.
"""

from __future__ import annotations

import os

from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _line(doc, text, size=13, bold=False, align="center", space_after=6,
          italic=False, indent=0.0):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.first_line_indent = Cm(indent)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    return p


def _page_border(section) -> None:
    """Khung viền đôi quanh trang bìa (bám theo lề văn bản, chừa gáy)."""
    sect_pr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "text")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "double")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "12")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    sect_pr.append(borders)


def _logo(doc, path: str, width_cm: float = 3.2) -> None:
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(path, width=Cm(width_cm))


def _new_cover_section(doc):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
    return sec


def add_cover(doc, info: dict) -> None:
    """Thêm ba trang bìa vào đầu tài liệu (đang ở section đầu tiên)."""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
    _page_border(sec)

    # ---------------- Trang 1: bìa chính ----------------
    _line(doc, "TRƯỜNG ĐẠI HỌC KỸ THUẬT CÔNG NGHIỆP", 13, True, space_after=2)
    _line(doc, info["khoa"], 13, True, space_after=2)
    _logo(doc, info["logo"])
    _line(doc, "", 10, space_after=30)
    _line(doc, "BÁO CÁO BÀI TẬP LỚN", 20, True, space_after=4)
    _line(doc, info["mon_hoc"], 18, True, space_after=40)
    _line(doc, f"NGÀNH    : {info['nganh']}", 16, True, space_after=6)
    _line(doc, f"HỆ          : {info['he']}", 16, True, space_after=46)
    _line(doc, f"ĐỀ TÀI: {info['de_tai']}", 16, True, space_after=60)
    _line(doc, f"THÁI NGUYÊN - {info['nam']}", 14, True, space_after=0)

    # ---------------- Trang 2: bìa lót ----------------
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _line(doc, "TRƯỜNG ĐẠI HỌC KỸ THUẬT CÔNG NGHIỆP", 13, True, space_after=2)
    _line(doc, info["khoa"], 13, True, space_after=2)
    _logo(doc, info["logo"])
    _line(doc, "", 10, space_after=24)
    _line(doc, "BÁO CÁO BÀI TẬP LỚN", 20, True, space_after=4)
    _line(doc, info["mon_hoc"], 18, True, space_after=4)
    _line(doc, f"NĂM HỌC: {info['nam_hoc']}", 14, True, space_after=4)
    _line(doc, f"BỘ MÔN: {info['bo_mon']}", 14, True, space_after=4)
    _line(doc, f"ĐỀ TÀI: {info['de_tai']}", 16, True, space_after=36)
    _line(doc, f"GIÁO VIÊN HƯỚNG DẪN   : {info['gvhd']}", 14, True, "left", space_after=8)
    _line(doc, f"HỌ VÀ TÊN SINH VIÊN     : {info['sv']}", 14, True, "left", space_after=8)
    _line(doc, f"LỚP                            : {info['lop']}", 14, True, "left",
          space_after=8)
    _line(doc, f"MSSV                        : {info['mssv']}", 14, True, "left", space_after=60)
    _line(doc, f"THÁI NGUYÊN - {info['nam']}", 14, True, space_after=0)

    # ---------------- Trang 3: nhiệm vụ bài tập lớn ----------------
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _line(doc, info["khoa"], 13, True, space_after=2)
    _line(doc, "BỘ MÔN: " + info["bo_mon"].upper(), 13, True, space_after=16)
    _line(doc, "NHIỆM VỤ BÀI TẬP LỚN", 16, True, space_after=16)
    _line(doc, f"Sinh viên: {info['sv']}          Lớp: {info['lop']}          "
               f"MSSV: {info['mssv']}", 13, space_after=6, align="left")
    _line(doc, f"Giáo viên hướng dẫn: {info['gvhd']}", 13, space_after=6, align="left")
    _line(doc, "1. Tên đề tài", 13, True, "left", space_after=2)
    _line(doc, info["de_tai"], 13, align="left", space_after=8, indent=0.6)
    _line(doc, "2. Các số liệu ban đầu", 13, True, "left", space_after=2)
    _line(doc, "Bộ dữ liệu lưới mẫu gồm sáu bản đồ có kích thước và mật độ vật cản khác "
               "nhau, trong đó có bản đồ địa hình mang trọng số chi phí di chuyển.", 13,
          align="left", space_after=8, indent=0.6)
    _line(doc, "3. Nội dung các phần thuyết minh và tính toán", 13, True, "left",
          space_after=2)
    for t in [
        "Phân tích bài toán tìm đường trên lưới, phát biểu mô hình đồ thị và hàm chi phí.",
        "Trình bày cơ sở lý thuyết của thuật toán A*, các hàm ước lượng h(n) và tính khả dụng.",
        "Chứng minh tính tối ưu của A* và phân tích độ phức tạp thời gian, bộ nhớ.",
        "Thiết kế chương trình gồm kiến trúc phân tầng, sơ đồ lớp và cấu trúc dữ liệu.",
        "Cài đặt chương trình có giao diện, kiểm thử và thực nghiệm đo đạc trên dữ liệu mẫu.",
    ]:
        _line(doc, "- " + t, 13, align="left", space_after=2, indent=0.6)
    _line(doc, "4. Các sản phẩm, kết quả", 13, True, "left", space_after=6)
    _line(doc, "5. Ngày giao nhiệm vụ: " + info.get("ngay_giao", ""), 13, align="left",
          space_after=6)
    _line(doc, "6. Ngày hoàn thành nhiệm vụ: " + info.get("ngay_ht", ""), 13, align="left",
          space_after=26)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    _line(doc, "BỘ MÔN (Ký, ghi rõ họ tên)          GIÁO VIÊN HƯỚNG DẪN (Ký, ghi rõ họ tên)",
          13, True, "center", space_after=0)


def add_body_section(doc):
    """Tạo section cho thân báo cáo: lề chuẩn học thuật, đánh số trang từ 1."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(3.0)
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.0)
    # Bỏ khung viền trang bìa kế thừa từ section trước (python-docx copy sectPr)
    for old in sec._sectPr.findall(qn("w:pgBorders")):
        sec._sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sec._sectPr.append(pg)
    sec.footer.is_linked_to_previous = False
    return sec
