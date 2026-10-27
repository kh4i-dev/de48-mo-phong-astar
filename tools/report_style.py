"""
report_style.py - Bộ dựng style học thuật cho báo cáo BTL (python-docx).

Chuẩn áp dụng (theo skill docx-report, mẫu TNUT):
  - Times New Roman 13 pt, căn đều hai bên, thụt đầu dòng 1 tab, giãn dòng 1.35
  - Heading 1: TNR đen đậm, ngắt trang trước; Heading 2/3/4 keep_with_next
  - Lề: trên 3.0, dưới 3.0, trái 3.5 (chừa gáy), phải 2.0 cm
  - Bảng: nền trắng 100%, tiêu đề cột in đậm căn trái, chống gãy hàng
  - Caption hình/bảng: đặt ngay dưới, chữ thường in nghiêng, căn giữa
"""

from __future__ import annotations

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
BODY_SIZE = 13
TABLE_SIZE = 11
CAPTION_SIZE = 12

NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


# ---------------------------------------------------------------------- #
# Khởi tạo
# ---------------------------------------------------------------------- #
def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(3.0)
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.0)

    # Đánh số trang bắt đầu lại từ 1 ở thân báo cáo (bìa ghép sau sẽ không đánh số)
    sect_pr = sec._sectPr
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    sect_pr.append(pg_num)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_SIZE)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after = Pt(2), Pt(4)
    pf.line_spacing = 1.35
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True

    _style_headings(doc)
    _ensure_caption_style(doc, "hình", CAPTION_SIZE)
    _ensure_caption_style(doc, "bảng", CAPTION_SIZE)
    _ensure_style(doc, "mã nguồn", "Consolas", 9.5, spacing=1.0)
    return doc


def _style_headings(doc: Document) -> None:
    sizes = {1: 16, 2: 14, 3: 13, 4: 12}
    for level, size in sizes.items():
        st = doc.styles[f"Heading {level}"]
        st.font.name = FONT
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(size)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        p = st.paragraph_format
        p.keep_with_next = True
        p.space_before = Pt(10 if level == 1 else 8)
        p.space_after = Pt(6)
        p.line_spacing = 1.2
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.page_break_before = (level == 1)


def _ensure_style(doc: Document, name: str, font: str, size: float,
                  italic: bool = False, spacing: float = 1.0) -> None:
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = font
    st.font.size = Pt(size)
    st.font.italic = italic
    st.font.bold = False
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.line_spacing = spacing
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)


def _ensure_caption_style(doc: Document, name: str, size: float) -> None:
    _ensure_style(doc, name, FONT, size, italic=True)
    st = doc.styles[name]
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_before = Pt(4)
    st.paragraph_format.space_after = Pt(10)
    st.paragraph_format.first_line_indent = Cm(0)


# ---------------------------------------------------------------------- #
# Khối nội dung
# ---------------------------------------------------------------------- #
def heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = FONT
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.italic = False
    if level == 1:
        h.paragraph_format.page_break_before = True
    h.paragraph_format.keep_with_next = True


def para(doc: Document, text: str, indent: bool = True, italic: bool = False,
         align: str = "justify") -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(BODY_SIZE)
    run.italic = italic
    p.paragraph_format.first_line_indent = Cm(1.0) if indent else Cm(0)
    p.paragraph_format.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        run.font.name = FONT
        run.font.size = Pt(BODY_SIZE)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)


def formula(doc: Document, expr: str, number: str = "") -> None:
    """Công thức căn giữa, số thứ tự sát lề phải (dùng tab dừng tường minh)."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    ppr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab_center = OxmlElement("w:tab")
    tab_center.set(qn("w:val"), "center")
    tab_center.set(qn("w:pos"), str(int(7.2 * 567)))
    tab_right = OxmlElement("w:tab")
    tab_right.set(qn("w:val"), "right")
    tab_right.set(qn("w:pos"), str(int(15.5 * 567)))
    tabs.append(tab_center)
    tabs.append(tab_right)
    ppr.append(tabs)
    run = p.add_run("\t" + expr + ("\t(" + number + ")" if number else ""))
    run.font.name = FONT
    run.font.size = Pt(BODY_SIZE)
    run.italic = True


def caption(doc: Document, text: str, style: str) -> None:
    p = doc.add_paragraph(text, style=style)
    for r in p.runs:
        r.font.name = FONT
        r.font.size = Pt(CAPTION_SIZE)
        r.font.italic = True
        r.font.bold = False
    # Triệt tiêu thụt lề kế thừa từ style Normal
    ppr = p._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLine"), "0")


def figure(doc: Document, image_path: str, caption_text: str,
           width_cm: float = 15.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(image_path, width=Cm(width_cm))
    caption(doc, caption_text, "hình")


def table(doc: Document, caption_text: str, headers: list[str], rows: list[list[str]],
          widths_cm: list[float], font_size: float = TABLE_SIZE,
          center_cols: tuple[int, ...] = ()) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    tbl_pr = t._tbl.tblPr
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)

    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        t._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)
    t._tbl.insert(t._tbl.index(tbl_pr) + 1, grid)

    hdr = t.rows[0]._tr.get_or_add_trPr()
    hdr.append(parse_xml(f"<w:tblHeader {NS}/>"))

    all_rows = [headers] + rows
    for ri, data in enumerate(all_rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(parse_xml(f"<w:cantSplit {NS}/>"))
        for ci, text in enumerate(data):
            cell = row.cells[ci]
            tc_pr = cell._tc.get_or_add_tcPr()
            for old in tc_pr.findall(qn("w:tcW")):
                tc_pr.remove(old)
            tc_pr.append(parse_xml(f'<w:tcW {NS} w:w="{int(widths_cm[ci] * 567)}" '
                                   f'w:type="dxa"/>'))
            for old_shd in tc_pr.findall(qn("w:shd")):
                tc_pr.remove(old_shd)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(text))
            run.font.name = FONT
            run.font.size = Pt(font_size)
            run.font.bold = (ri == 0)
            run.font.italic = False
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif ci in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ppr = p._p.get_or_add_pPr()
            ind = ppr.find(qn("w:ind"))
            if ind is None:
                ind = OxmlElement("w:ind")
                ppr.append(ind)
            ind.set(qn("w:firstLine"), "0")
            ind.set(qn("w:left"), "0")
            ind.set(qn("w:right"), "0")
    caption(doc, caption_text, "bảng")


def code_block(doc: Document, code: str, max_lines: int = 40) -> None:
    lines = code.rstrip("\n").split("\n")[:max_lines]
    for line in lines:
        # Giữ thụt lề mã nguồn: đổi khoảng trắng đầu dòng thành khoảng trắng không ngắt
        stripped = line.lstrip(" ")
        indent = len(line) - len(stripped)
        line = "\u00a0" * indent + stripped
        p = doc.add_paragraph(style="mã nguồn")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line.strip() else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------- #
# Field tự động: MỤC LỤC / DANH MỤC HÌNH / DANH MỤC BẢNG / SỐ TRANG
# ---------------------------------------------------------------------- #
def add_toc_field(doc: Document, instr: str = r'TOC \o "1-3" \h \z \u',
                  placeholder: str = "Mục lục sẽ hiện sau khi cập nhật trường "
                                     "(Ctrl+A rồi F9 trong Word)") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(fld)
    r._r.append(instr_el)
    r._r.append(sep)
    r._r.append(t)
    r._r.append(end)


def add_page_numbers(doc: Document, sections=None) -> None:
    for sec in (sections if sections is not None else doc.sections):
        footer = sec.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "begin")
        instr_el = OxmlElement("w:instrText")
        instr_el.text = "PAGE"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(fld)
        run._r.append(instr_el)
        run._r.append(end)
        run.font.name = FONT
        run.font.size = Pt(11)


def word_count(doc: Document) -> int:
    n = 0
    for p in doc.paragraphs:
        n += len(p.text.split())
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                n += len(c.text.split())
    return n


def save(doc: Document, path: str) -> str:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)
    return path
