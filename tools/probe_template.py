"""Đọc template báo cáo KTLTNC có sẵn: bìa, style, lề, font, caption."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Pt

path = sys.argv[1]
doc = Document(path)

print("=" * 70)
print("SECTION / TRANG")
for i, s in enumerate(doc.sections):
    print(f"  section {i}: {s.page_width.cm:.2f} x {s.page_height.cm:.2f} cm | "
          f"lề T{s.top_margin.cm:.2f} B{s.bottom_margin.cm:.2f} "
          f"L{s.left_margin.cm:.2f} R{s.right_margin.cm:.2f} | "
          f"header_dist={s.header_distance.cm:.2f}")
    for hp in s.header.paragraphs:
        if hp.text.strip():
            print(f"     header: {hp.text!r}")
    for fp in s.footer.paragraphs:
        if fp.text.strip():
            print(f"     footer: {fp.text!r}")

print("=" * 70)
print("STYLE ĐANG DÙNG")
seen = set()
for p in doc.paragraphs:
    n = p.style.name
    if n not in seen:
        seen.add(n)
        st = doc.styles[n] if n in [s.name for s in doc.styles] else None
        try:
            f = st.font
            info = (f"{f.name or '-'} / {f.size.pt if f.size else '-'}pt / "
                    f"bold={f.bold} italic={f.italic}")
        except Exception as e:
            info = f"?? {e}"
        print(f"  - {n:26} {info}")

print("=" * 70)
print("70 ĐOẠN ĐẦU (bìa + front matter)")
for i, p in enumerate(doc.paragraphs[:70]):
    t = p.text.strip()
    if not t:
        continue
    runs = p.runs[0] if p.runs else None
    fn = (runs.font.name if runs else None) or "-"
    sz = (runs.font.size.pt if runs and runs.font.size else None) or "-"
    print(f"  [{i:3}] {p.style.name:22} | {fn:16} {str(sz):>5} | {t[:95]}")

print("=" * 70)
print("BẢNG: số lượng", len(doc.tables))
for ti, t in enumerate(doc.tables[:3]):
    print(f"  bảng {ti}: {len(t.rows)}x{len(t.columns)} | "
          f"header={[c.text[:18] for c in t.rows[0].cells]}")
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                print(f"     header run: {r.text[:20]!r} bold={r.font.bold} "
                      f"align={p.alignment} size={r.font.size.pt if r.font.size else None}")
            break

print("=" * 70)
print("CAPTION (đoạn chứa 'Hình'/'Bảng' ở đầu)")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith(("Hình", "Bảng")) and len(t) < 120:
        r = p.runs[0] if p.runs else None
        print(f"  [{i:3}] style={p.style.name:16} italic={r.font.italic if r else '-'} "
              f"align={p.alignment} size={r.font.size.pt if r and r.font.size else '-'} | {t[:80]}")

print("=" * 70)
print("INLINE SHAPES:", len(doc.inline_shapes))
